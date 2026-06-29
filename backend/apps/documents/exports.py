import io
import os
from copy import copy
from django.conf import settings
from django.http import HttpResponse
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
import base64


TEMPLATE_PATH = os.path.join(settings.BASE_DIR, 'apps', 'documents', 'templates_xlsx', 'business_permit.xlsx')


def _fill_labeled_cell(ws, coord, label, value):
    """
    Many cells in this template are formatted as "Label: ______" where the
    underscores are a printed blank meant to be replaced with the real
    answer. This rewrites that cell as "Label: <value>", preserving the
    original font/border/alignment of the cell exactly.
    """
    cell = ws[coord]
    original_font = copy(cell.font)
    original_border = copy(cell.border)
    original_alignment = copy(cell.alignment)
    original_fill = copy(cell.fill)

    cell.value = f"{label} {value}" if value else cell.value
    cell.font = original_font
    cell.border = original_border
    cell.alignment = original_alignment
    cell.fill = original_fill


def fill_business_permit_front(ws, fd):
    """
    fd is the document's form_data dict. Keys map directly to the React
    multi-step form's field names — see BusinessPermitWizard.jsx.
    """
    _fill_labeled_cell(ws, 'B2', 'TAX YEAR', fd.get('tax_year', ''))

    _fill_labeled_cell(ws, 'B11', 'Date of Application:', fd.get('date_of_application', ''))
    _fill_labeled_cell(ws, 'B12', 'TIN No.:', fd.get('tin_no', ''))
    _fill_labeled_cell(ws, 'E11', 'DTI/SEC/CDA Registration No.:', fd.get('dti_reg_no', ''))
    _fill_labeled_cell(ws, 'E12', 'DTI/SEC/CDA Registration Date:', fd.get('dti_reg_date', ''))

    _fill_labeled_cell(ws, 'B18', 'Last Name:', fd.get('last_name', ''))
    _fill_labeled_cell(ws, 'D18', 'First Name:', fd.get('first_name', ''))
    _fill_labeled_cell(ws, 'F18', 'Middle Name:', fd.get('middle_name', ''))
    _fill_labeled_cell(ws, 'B19', 'Business Name:', fd.get('business_name', ''))
    _fill_labeled_cell(ws, 'B20', 'Trade name / Franchise:', fd.get('trade_name', ''))

    _fill_labeled_cell(ws, 'B22', 'Business Address:', fd.get('business_address', ''))
    _fill_labeled_cell(ws, 'B23', 'Postal Code:', fd.get('business_postal_code', ''))
    _fill_labeled_cell(ws, 'E23', 'Email Address:', fd.get('business_email', ''))
    _fill_labeled_cell(ws, 'B24', 'Telephone No.:', fd.get('business_telephone', ''))
    _fill_labeled_cell(ws, 'E24', 'Mobile No.:', fd.get('business_mobile', ''))

    _fill_labeled_cell(ws, 'B25', "Owner's Address:", fd.get('owner_address', ''))
    _fill_labeled_cell(ws, 'B26', 'Postal Code:', fd.get('owner_postal_code', ''))
    _fill_labeled_cell(ws, 'E26', 'Email Address:', fd.get('owner_email', ''))
    _fill_labeled_cell(ws, 'B27', 'Telephone No.:', fd.get('owner_telephone', ''))
    _fill_labeled_cell(ws, 'E27', 'Mobile No.:', fd.get('owner_mobile', ''))

    _fill_labeled_cell(ws, 'B28', 'In case of emergency, provide name of contact person:', fd.get('emergency_contact_name', ''))
    _fill_labeled_cell(ws, 'B29', 'Telephone/Mobile No.:', fd.get('emergency_contact_phone', ''))
    _fill_labeled_cell(ws, 'E29', 'Address:', fd.get('emergency_contact_address', ''))

    _fill_labeled_cell(ws, 'B30', 'Business Area (in sq m.):', fd.get('floor_area', ''))
    _fill_labeled_cell(ws, 'D30', 'Total No. Employees in Establishment:', fd.get('total_employees', ''))
    _fill_labeled_cell(ws, 'D31', 'Male:', fd.get('employees_male', ''))
    _fill_labeled_cell(ws, 'E31', 'Female:', fd.get('employees_female', ''))
    _fill_labeled_cell(ws, 'F31', 'LGU:', fd.get('employees_residing_lgu', ''))

    if fd.get('is_rented'):
        _fill_labeled_cell(ws, 'B33', "Lessor's Full Name:", fd.get('lessor_name', ''))
        _fill_labeled_cell(ws, 'B34', "Lessor's Full Address:", fd.get('lessor_address', ''))
        _fill_labeled_cell(ws, 'B35', "Lessor's Full Telephone/Mobile No.:", fd.get('lessor_phone', ''))
        _fill_labeled_cell(ws, 'F35', 'Monthly Rental:', fd.get('monthly_rental', ''))
        _fill_labeled_cell(ws, 'B36', 'Name of the Building Rented:', fd.get('building_name', ''))
        _fill_labeled_cell(ws, 'B37', 'Address of the Building being Rented:', fd.get('building_address', ''))

    _fill_labeled_cell(ws, 'E54', 'POSITION/TITLE', fd.get('position_title', ''))



def fill_business_permit_back(ws, fd):
    _fill_labeled_cell(ws, 'B20', 'DATE:', fd.get('date_of_application', ''))
    _fill_labeled_cell(ws, 'B21', 'APPLICATION NO.:', fd.get('application_no', ''))
    _fill_labeled_cell(ws, 'B23', 'Name of Applicant /  Owner :', fd.get('applicant_name', ''))
    _fill_labeled_cell(ws, 'B24', 'Name of Business :', fd.get('business_name', ''))
    _fill_labeled_cell(ws, 'B25', f"Total Floor Area: {fd.get('floor_area', '')}  Contact No.:", fd.get('contact_no', ''))
    _fill_labeled_cell(ws, 'B26', 'Address of Establishment:', fd.get('address', ''))

def insert_qr_block(ws, document):
    """
    Replaces the decorative logo image in the top-right corner of the
    form (originally a 'Gilas Sariaya' graphic unrelated to this system)
    with the document's QR code, so it's visible right at the top when
    printed — exactly where a stamp/logo would normally sit.
    """
    # Remove the existing top-right logo image (col=6, row=0 in the
    # original template) without touching the official seal (top-left).
    images_to_keep = []
    for img in ws._images:
        try:
            col = img.anchor._from.col
            row = img.anchor._from.row
        except Exception:
            images_to_keep.append(img)
            continue
        if col == 6 and row == 0:
            continue  # this is the logo we're replacing — skip it
        images_to_keep.append(img)
    ws._images = images_to_keep

    if document.qr_code:
        try:
            img = XLImage(document.qr_code.path)
            img.width = 130
            img.height = 130
            img.anchor = 'G1'
            ws.add_image(img)
        except Exception:
            pass


def insert_signature(ws, fd):
    """
    The applicant draws their signature on a canvas in the browser, which
    gets sent as a base64 PNG data URL in form_data['signature']. This
    decodes it and embeds it directly above the
    'SIGNATURE OF APPLICANT/TAXPAYER OVER PRINTED NAME' line.
    """
    signature_data = fd.get('signature')
    if not signature_data:
        return
    try:
        header, encoded = signature_data.split(',', 1)
        img_bytes = base64.b64decode(encoded)
        img_buffer = io.BytesIO(img_bytes)
        img = XLImage(img_buffer)
        img.width = 180
        img.height = 50
        img.anchor = 'E49'
        ws.add_image(img)
    except Exception:
        pass  # if signature data is malformed, export still proceeds without it


def export_document_as_xlsx(document):
    if document.document_type.name == 'Business Permit' and os.path.exists(TEMPLATE_PATH):
        wb = load_workbook(TEMPLATE_PATH)
        fd = document.form_data or {}

        front = wb['FRONT']
        back = wb['BACK']

        fill_business_permit_front(front, fd)
        fill_business_permit_back(back, fd)
        insert_qr_block(front, document)
        insert_signature(front, fd)
    else:
        # Fallback for any document type without a matching template yet.
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws['A1'] = f"{document.document_type.name} — {document.tracking_number}"
        if document.qr_code:
            try:
                img = XLImage(document.qr_code.path)
                img.anchor = 'A3'
                ws.add_image(img)
            except Exception:
                pass
        row = 10
        for key, value in (document.form_data or {}).items():
            ws.cell(row=row, column=1, value=key.replace('_', ' ').title())
            ws.cell(row=row, column=2, value=value)
            row += 1

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{document.tracking_number}.xlsx"'
    return response